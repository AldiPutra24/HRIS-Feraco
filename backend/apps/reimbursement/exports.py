"""Reimbursement XLSX export (openpyxl), same approach as payroll/exports.py.

Rows come from the caller's already-scoped queryset, so the file never holds
more than the user may see in the list (HR: all; others: own only).
"""
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Reimbursement

RUPIAH = '"Rp" #,##0'
DATE = 'DD/MM/YYYY'

# (header, width, number_format)
COLUMNS = [
    ('Karyawan', 28, None),
    ('Kategori', 20, None),
    ('Kategori Project', 26, None),
    ('Deskripsi', 40, None),
    ('Tanggal', 13, DATE),
    ('Nominal Diajukan', 18, RUPIAH),
    ('Nominal Disetujui', 18, RUPIAH),
    ('Status', 18, None),
    ('Nama Bank', 16, None),
    ('Nama Pemilik Rekening', 26, None),
    ('Nomor Rekening', 20, None),
    ('Email', 28, None),
    ('Reviewer HR Staff', 22, None),
    ('Approver HR Lead', 22, None),
    ('Tanggal Pembayaran', 18, DATE),
]


def _user_name(user):
    if user is None:
        return ''
    personnel = getattr(user, 'personnel', None)
    name = getattr(personnel, 'full_name', '') if personnel else ''
    return name or user.get_full_name() or user.get_username()


def _project(obj):
    if obj.project_category == 'OTHER':
        return obj.project_category_other or 'Other'
    return dict(Reimbursement.PROJECT_CATEGORY_CHOICES).get(obj.project_category, obj.project_category or '')


def _local_date(dt):
    return timezone.localtime(dt).date() if dt else None


def build_reimbursement_xlsx(queryset) -> HttpResponse:
    rows = queryset.select_related(
        'employee', 'category', 'amount_set_by', 'reviewer',
    ).order_by('-transaction_date', '-id')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Reimbursement'
    ws.append([c[0] for c in COLUMNS])
    header_fill = PatternFill('solid', fgColor='E8EEF7')
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')

    for obj in rows:
        ws.append([
            obj.employee.full_name,
            obj.category.name if obj.category_id else '',
            _project(obj),
            obj.description or '',
            obj.transaction_date,
            float(obj.amount) if obj.amount is not None else None,
            float(obj.approved_amount) if obj.approved_amount is not None else None,
            obj.get_status_display(),
            obj.bank_name or '',
            obj.bank_account_name or '',
            # Text, so leading zeros and long numbers are kept exactly.
            obj.bank_account_number or '',
            obj.contact_email or '',
            _user_name(obj.amount_set_by),
            # `reviewer` is the final decision maker; it is the HR Lead
            # approver only when the request was actually approved.
            _user_name(obj.reviewer) if obj.approved_at else '',
            _local_date(obj.paid_at),
        ])

    for idx, (_, width, fmt) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = fmt
    for (cell,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for (cell,) in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        cell.number_format = '@'

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="reimbursement-{timezone.localdate().isoformat()}.xlsx"'
    )
    wb.save(response)
    return response
